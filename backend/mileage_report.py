"""Builds the HR-format mileage workbook from logged trips.

The styled shell ships as a vendored template (assets/mileage_report_template.xlsx
— HR's 'Mileage Tracker 2026' with the approved rework baked in: no odometer
columns, ROUTE FROM/TO address columns, no frozen panes, landscape print at a
true 100% scale, EMP ID / DEPT left blank for hand-writing). This module only
fills the name, the rows, and the totals; every visual choice lives in the
template so the download matches the sample Ethan approved.
"""
import io
import math
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

_TEMPLATE = Path(__file__).parent / "assets" / "mileage_report_template.xlsx"
_SHEET = "Business Vehicle Mileage Log"
_START = 10       # first data row of the template grid
_GRID_END = 38    # last pre-styled grid row; past it we copy row-10 styles
_TOTAL_PROTO = 39  # the template's (emptied) grand-total row — style source


def _fmt_rate(rate: float) -> str:
    """$0.70 and $0.725 — 3 decimals only when they matter."""
    s = f"{float(rate):.3f}"
    return f"${s[:-1] if s.endswith('0') else s}"


def _lines(text: str, width: int) -> int:
    """Rough wrapped-line count for 10pt Century Gothic in a `width`-unit
    column — drives explicit row heights, since openpyxl-written files never
    get Excel's auto-fit and clipped rows were the #1 print complaint."""
    if not text:
        return 1
    return max(1, math.ceil(len(text) / max(width - 3, 8)))


def build_report(user_name: str, days: list, header_rate: float | None,
                 total_dollars: float) -> bytes:
    """days: [{date: 'YYYY-MM-DD', rows: [{purpose, frm, to, miles}], rate}]
    in report order. header_rate is the single rate shared by every trip, or
    None when rates vary — then the header reads VARIES, the reimbursement is
    the per-trip dollar sum, and each DAY TOTAL row notes its own rate."""
    wb = load_workbook(_TEMPLATE)
    ws = wb[_SHEET]
    ws["C4"] = user_name or None

    proto_row = {c: copy(ws.cell(row=_START, column=c)._style) for c in range(2, 8)}
    proto_total = {c: copy(ws.cell(row=_TOTAL_PROTO, column=c)._style) for c in range(2, 8)}

    def styled(r, c, value=None):
        cell = ws.cell(row=r, column=c)
        if r > _GRID_END:
            cell._style = copy(proto_row[c])
        if value is not None:
            cell.value = value
        return cell

    def bolden(cell):
        fnt = copy(cell.font)
        fnt.bold = True
        cell.font = fnt

    mixed = header_rate is None
    r = _START
    last_day_total = _START
    for day in days:
        d = datetime.strptime(day["date"], "%Y-%m-%d")
        first = r
        for row in day["rows"]:
            styled(r, 2, d).number_format = "m/d/yyyy"
            styled(r, 3, row["purpose"] or None)
            styled(r, 4, row["frm"])
            styled(r, 5, row["to"])
            styled(r, 6, round(float(row["miles"]), 1)).number_format = "0.0"
            styled(r, 7)  # keeps the comments column's grid style past row 38
            n = max(_lines(row["purpose"], 24), _lines(row["frm"], 27),
                    _lines(row["to"], 27))
            ws.row_dimensions[r].height = 6 + n * 13.5
            r += 1
        styled(r, 2, d).number_format = "m/d/yyyy"
        lab = styled(r, 5, "DAY TOTAL")
        a = copy(lab.alignment)
        a.horizontal = "right"
        lab.alignment = a
        bolden(lab)
        tot = styled(r, 6, f"=SUM(F{first}:F{r - 1})")
        tot.number_format = "0.0"
        bolden(tot)
        if mixed:
            styled(r, 7, f"@ {_fmt_rate(day['rate'])}/mi")
        ws.row_dimensions[r].height = 18
        last_day_total = r
        r += 2  # blank spacer row between days

    # Grand total: SUMIF over the DAY TOTAL rows only, so leg rows are never
    # double-counted. Position floats with the data; styles come from the
    # template's emptied row 39.
    total_row = r
    for c in range(2, 8):
        ws.cell(row=total_row, column=c)._style = copy(proto_total[c])
    ws.cell(row=total_row, column=2, value="Total")
    g = ws.cell(
        row=total_row, column=6,
        value=f'=SUMIF(E{_START}:E{last_day_total},"DAY TOTAL",F{_START}:F{last_day_total})',
    )
    g.number_format = "0.0"
    ws.row_dimensions[total_row].height = 18

    ws["G5"] = f"=F{total_row}"
    ws["G5"].number_format = "0.0"
    if mixed:
        ws["G4"] = "VARIES"
        ws["G6"] = round(total_dollars, 2)  # per-trip sum; no single rate exists
    else:
        ws["G4"] = round(header_rate, 4)
        ws["G6"] = "=G4*G5"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
